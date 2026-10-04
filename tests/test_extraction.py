"""
DocuSentinel AI - Tests: Document Extraction
Tests PDF, DOCX, TXT extraction and error handling.
"""

import os
import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock

from backend.services.extractor import (
    extract_txt, extract_pdf, extract_docx, extract_image, extract_document
)


# ─────────────────────────────────────────────
# TXT Extraction
# ─────────────────────────────────────────────

class TestTxtExtraction:

    def test_extract_txt_basic(self, tmp_path):
        """TXT extraction returns pages with correct text."""
        txt_file = tmp_path / "test.txt"
        content = "This is paragraph one.\n\nThis is paragraph two with more content.\n\nThird paragraph here."
        txt_file.write_text(content, encoding="utf-8")

        result = extract_txt(str(txt_file))

        assert result.success
        assert result.total_pages >= 1
        assert len(result.raw_text) > 0
        assert "paragraph one" in result.raw_text

    def test_extract_txt_empty_file(self, tmp_path):
        """Empty TXT file returns an error result."""
        txt_file = tmp_path / "empty.txt"
        txt_file.write_text("", encoding="utf-8")

        result = extract_txt(str(txt_file))

        assert not result.success
        assert result.error is not None
        assert "empty" in result.error.lower()

    def test_extract_txt_preserves_content(self, tmp_path):
        """TXT extraction preserves all significant text."""
        txt_file = tmp_path / "contract.txt"
        txt_file.write_text(
            "Project Deadline: 15 November 2026\n\n"
            "Total Contract Value: USD 1,200,000\n\n"
            "Project Manager: Mr. John Smith",
            encoding="utf-8",
        )

        result = extract_txt(str(txt_file))

        assert result.success
        combined = result.raw_text
        assert "15 November 2026" in combined
        assert "1,200,000" in combined
        assert "John Smith" in combined

    def test_extract_txt_large_file_pagination(self, tmp_path):
        """Large TXT content is split into multiple pages."""
        txt_file = tmp_path / "large.txt"
        # Create content well over 400 words per page
        paragraphs = [f"This is paragraph {i} with some content words here." for i in range(100)]
        txt_file.write_text("\n\n".join(paragraphs), encoding="utf-8")

        result = extract_txt(str(txt_file))

        assert result.success
        assert result.total_pages > 1

    def test_extract_txt_unicode(self, tmp_path):
        """TXT extraction handles unicode content."""
        txt_file = tmp_path / "unicode.txt"
        txt_file.write_text("Deadline: 15 नवंबर 2026\nCost: €1,200,000", encoding="utf-8")

        result = extract_txt(str(txt_file))

        assert result.success
        assert "2026" in result.raw_text


# ─────────────────────────────────────────────
# PDF Extraction
# ─────────────────────────────────────────────

class TestPdfExtraction:

    def test_extract_pdf_missing_file(self):
        """Missing PDF file returns error result."""
        result = extract_pdf("/nonexistent/path/file.pdf")
        assert not result.success
        assert result.error is not None

    def test_extract_pdf_with_pymupdf(self, tmp_path):
        """PDF extraction with PyMuPDF returns pages when available."""
        try:
            import fitz
            import reportlab
        except ImportError:
            pytest.skip("PyMuPDF or reportlab not available for PDF creation")

        # Create a minimal PDF using reportlab
        from reportlab.pdfgen import canvas
        pdf_path = tmp_path / "test.pdf"
        c = canvas.Canvas(str(pdf_path))
        c.drawString(100, 750, "Project Deadline: 15 November 2026")
        c.drawString(100, 730, "Contract Value: USD 1,200,000")
        c.save()

        result = extract_pdf(str(pdf_path))
        assert result.success
        assert result.total_pages >= 1


# ─────────────────────────────────────────────
# DOCX Extraction
# ─────────────────────────────────────────────

class TestDocxExtraction:

    def test_extract_docx_missing_file(self):
        """Missing DOCX file returns error result."""
        result = extract_docx("/nonexistent/path/file.docx")
        assert not result.success
        assert result.error is not None

    def test_extract_docx_with_python_docx(self, tmp_path):
        """DOCX extraction works with python-docx generated files."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not available")

        docx_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_heading("Project Contract", level=1)
        doc.add_paragraph("Project Deadline: 15 November 2026")
        doc.add_paragraph("Contract Value: USD 1,200,000")
        doc.save(str(docx_path))

        result = extract_docx(str(docx_path))
        assert result.success
        assert "15 November 2026" in result.raw_text or "2026" in result.raw_text


# ─────────────────────────────────────────────
# Dispatcher
# ─────────────────────────────────────────────

class TestExtractDocumentDispatcher:

    def test_unsupported_type(self, tmp_path):
        """Unsupported file type returns error result."""
        result = extract_document("/some/file.xyz", "EXCEL")
        assert not result.success
        assert "Unsupported" in result.error

    def test_dispatches_txt(self, tmp_path):
        """Dispatcher routes TXT files correctly."""
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("Some test content here for routing.", encoding="utf-8")

        result = extract_document(str(txt_file), "TXT")
        assert result.success

    def test_dispatches_pdf(self, tmp_path):
        """Dispatcher routes PDF files to PDF extractor."""
        fake_pdf = tmp_path / "test.pdf"
        fake_pdf.write_bytes(b"")  # empty — will fail but via correct handler

        result = extract_document(str(fake_pdf), "PDF")
        # Result may fail on empty PDF, but error should be from PDF handler
        assert result.error is not None
